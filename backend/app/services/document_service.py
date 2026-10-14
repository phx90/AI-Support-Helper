import io
import re
import uuid
from datetime import UTC, datetime
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import delete
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.core.exceptions import DocumentProcessingError, NotFoundError
from app.models.entities import Document, DocumentChunk, DocumentStatus
from app.rag.chunking import TechnicalTextChunker
from app.rag.embeddings import get_embedding_provider

ALLOWED = {".pdf", ".txt", ".md", ".docx"}
UPLOAD_DIR = Path("data/uploads")


def safe_filename(name: str) -> str:
    name = Path(name).name
    return re.sub(r"[^A-Za-z0-9._-]", "_", name)


def extract_pages(data: bytes, suffix: str) -> list[tuple[int | None, str]]:
    try:
        if suffix == ".pdf":
            return [
                (i + 1, p.extract_text() or "")
                for i, p in enumerate(PdfReader(io.BytesIO(data)).pages)
            ]
        if suffix == ".docx":
            return [(None, "\n".join(p.text for p in DocxDocument(io.BytesIO(data)).paragraphs))]
        return [(None, data.decode("utf-8"))]
    except Exception as exc:
        raise DocumentProcessingError("Could not extract document text") from exc


def validate_upload(filename: str, data: bytes) -> tuple[str, str]:
    name = safe_filename(filename)
    suffix = Path(name).suffix.lower()
    s = get_settings()
    if suffix not in ALLOWED:
        raise DocumentProcessingError("Unsupported file type. Use PDF, TXT, MD or DOCX.")
    if not data:
        raise DocumentProcessingError("The uploaded file is empty.")
    if len(data) > s.max_upload_mb * 1024 * 1024:
        raise DocumentProcessingError(f"File exceeds {s.max_upload_mb} MB limit.")
    return name, suffix


def create_document(db: Session, filename: str, data: bytes) -> Document:
    name, suffix = validate_upload(filename, data)
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(name=name, file_type=suffix[1:], size=len(data), storage_path="")
    db.add(doc)
    db.flush()
    path = UPLOAD_DIR / f"{doc.id}{suffix}"
    path.write_bytes(data)
    doc.storage_path = str(path)
    db.commit()
    db.refresh(doc)
    return doc


def process_document(db: Session, document_id: uuid.UUID) -> Document:
    doc = db.get(Document, document_id)
    if not doc:
        raise NotFoundError("Document not found")
    doc.status = DocumentStatus.processing
    doc.error_message = None
    db.commit()
    try:
        data = Path(doc.storage_path).read_bytes()
        chunks = TechnicalTextChunker(
            get_settings().chunk_size, get_settings().chunk_overlap
        ).split(extract_pages(data, "." + doc.file_type))
        if not chunks:
            raise DocumentProcessingError("No readable text found in document")
        embeddings = get_embedding_provider().embed([c.content for c in chunks])
        db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))
        for c, e in zip(chunks, embeddings, strict=True):
            db.add(
                DocumentChunk(
                    document_id=doc.id,
                    content=c.content,
                    chunk_index=c.chunk_index,
                    page=c.page,
                    section=c.section,
                    metadata_={
                        "document_id": str(doc.id),
                        "filename": doc.name,
                        "page": c.page,
                        "chunk_index": c.chunk_index,
                        "section": c.section,
                        "source": doc.name,
                    },
                    embedding=e,
                )
            )
        doc.status = DocumentStatus.indexed
        doc.chunk_count = len(chunks)
        doc.processed_at = datetime.now(UTC)
        db.commit()
        db.refresh(doc)
    except Exception as exc:
        db.rollback()
        doc = db.get(Document, document_id)
        doc.status = DocumentStatus.error
        doc.error_message = str(exc)[:1000]
        db.commit()
        raise
    return doc


def delete_document(db: Session, doc: Document) -> None:
    path = Path(doc.storage_path)
    db.delete(doc)
    db.commit()
    path.unlink(missing_ok=True)
